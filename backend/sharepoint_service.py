"""Read-only SharePoint access through Microsoft Graph."""

import os
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path, PurePosixPath
from urllib.parse import quote, urlparse

import httpx
import msal
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


BASE_DIR = Path(__file__).resolve().parent
MAX_SEARCH_RESULTS = 200
MAX_TEXT_BYTES = 2 * 1024 * 1024
MAX_DOWNLOAD_BYTES = 50 * 1024 * 1024
MAX_EXCEL_BYTES = 20 * 1024 * 1024
MAX_EXCEL_ROWS = 200
MAX_EXCEL_CELLS = 5000
TEXT_EXTENSIONS = {".txt", ".csv", ".json", ".md", ".xml", ".log"}


class SharePointError(RuntimeError):
    """A safe, user-facing SharePoint operation error."""


@dataclass(frozen=True)
class SharePointConfig:
    client_id: str
    tenant_id: str
    site_url: str
    download_dir: Path
    token_cache_path: Path = BASE_DIR / ".sharepoint_token_cache.json"

    @classmethod
    def from_env(cls) -> "SharePointConfig":
        client_id = os.getenv("SHAREPOINT_CLIENT_ID", "").strip()
        tenant_id = os.getenv("SHAREPOINT_TENANT_ID", "").strip()
        site_url = os.getenv("SHAREPOINT_SITE_URL","").strip()
        missing = [
            name
            for name, value in (
                ("SHAREPOINT_CLIENT_ID", client_id),
                ("SHAREPOINT_TENANT_ID", tenant_id),
            )
            if not value
        ]
        if missing:
            raise SharePointError(
                "Missing SharePoint configuration: " + ", ".join(missing)
            )
        download_dir = Path(
            os.getenv(
                "SHAREPOINT_DOWNLOAD_DIR",
                str(BASE_DIR / "downloads" / "sharepoint"),
            )
        )
        return cls(
            client_id=client_id,
            tenant_id=tenant_id,
            site_url=site_url,
            download_dir=download_dir,
        )

    @property
    def hostname(self) -> str:
        parsed = urlparse(self.site_url)
        if parsed.scheme != "https" or not parsed.hostname:
            raise SharePointError("SHAREPOINT_SITE_URL must be an HTTPS URL")
        return parsed.hostname

    @property
    def site_path(self) -> str:
        path = urlparse(self.site_url).path.rstrip("/")
        if not path:
            raise SharePointError("SHAREPOINT_SITE_URL must include a site path")
        return path


class SharePointService:
    SCOPES = ["Files.Read.All", "Sites.Read.All"]
    GRAPH_ROOT = "https://graph.microsoft.com/v1.0"

    def __init__(
        self,
        config: SharePointConfig | None = None,
        http_client: httpx.Client | None = None,
        auth_app=None,
        sleep=time.sleep,
    ):
        self.config = config or SharePointConfig.from_env()
        self.http = http_client or httpx.Client(timeout=30.0, follow_redirects=True)
        self.sleep = sleep
        self._token_cache = msal.SerializableTokenCache()
        if self.config.token_cache_path.exists():
            self._token_cache.deserialize(
                self.config.token_cache_path.read_text(encoding="utf-8")
            )
        self.auth_app = auth_app or msal.PublicClientApplication(
            self.config.client_id,
            authority=f"https://login.microsoftonline.com/{self.config.tenant_id}",
            token_cache=self._token_cache,
        )
        self._ids: tuple[str, str] | None = None

    def _get_access_token(self) -> str:
        accounts = self.auth_app.get_accounts()
        result = (
            self.auth_app.acquire_token_silent(self.SCOPES, account=accounts[0])
            if accounts
            else None
        )
        if not result:
            try:
                result = self.auth_app.acquire_token_interactive(scopes=self.SCOPES)
            except Exception:
                flow = self.auth_app.initiate_device_flow(scopes=self.SCOPES)
                if "user_code" not in flow:
                    raise SharePointError("Could not start Microsoft device login")
                print(flow["message"], file=sys.stderr, flush=True)
                result = self.auth_app.acquire_token_by_device_flow(flow)
        self._save_cache()
        if not result or "access_token" not in result:
            message = (result or {}).get(
                "error_description", "Microsoft sign-in failed"
            )
            raise SharePointError(message)
        return result["access_token"]

    def _save_cache(self) -> None:
        if self._token_cache.has_state_changed:
            self.config.token_cache_path.parent.mkdir(parents=True, exist_ok=True)
            self.config.token_cache_path.write_text(
                self._token_cache.serialize(), encoding="utf-8"
            )

    def _request(self, method: str, endpoint: str, **kwargs) -> httpx.Response:
        headers = dict(kwargs.pop("headers", {}))
        headers["Authorization"] = f"Bearer {self._get_access_token()}"
        response = None
        for attempt in range(3):
            try:
                response = self.http.request(
                    method,
                    self.GRAPH_ROOT + endpoint,
                    headers=headers,
                    **kwargs,
                )
            except httpx.HTTPError as exc:
                if attempt == 2:
                    raise SharePointError(
                        f"Microsoft Graph request failed: {exc}"
                    ) from exc
                self.sleep(1.0)
                continue
            if response.status_code not in {429, 500, 502, 503, 504}:
                break
            if attempt < 2:
                try:
                    delay = float(response.headers.get("Retry-After", 1))
                except ValueError:
                    delay = 1.0
                self.sleep(min(delay, 5.0))
        if response is None:
            raise SharePointError("Microsoft Graph request did not run")
        if response.status_code >= 400:
            try:
                message = response.json()["error"]["message"]
            except (KeyError, TypeError, ValueError):
                message = response.text or "Microsoft Graph request failed"
            raise SharePointError(
                f"Microsoft Graph {response.status_code}: {message}"
            )
        return response

    def _site_and_drive_ids(self) -> tuple[str, str]:
        if self._ids is None:
            site = self._request(
                "GET", f"/sites/{self.config.hostname}:{self.config.site_path}"
            ).json()
            drive = self._request("GET", f"/sites/{site['id']}/drive").json()
            self._ids = (site["id"], drive["id"])
        return self._ids

    @staticmethod
    def _clean_path(value: str) -> str:
        candidate = PurePosixPath(value.replace("\\", "/").strip())
        if candidate.is_absolute() or ".." in candidate.parts:
            raise SharePointError(
                "SharePoint paths must stay inside the configured site"
            )
        return "" if str(candidate) == "." else candidate.as_posix()

    def _item_endpoint(self, drive_id: str, path: str, suffix: str) -> str:
        cleaned = self._clean_path(path)
        if not cleaned:
            return f"/drives/{drive_id}/root/{suffix.lstrip('/')}"
        encoded = quote(cleaned, safe="/")
        return f"/drives/{drive_id}/root:/{encoded}:/{suffix.lstrip('/')}"

    @staticmethod
    def _normalize_item(item: dict) -> dict:
        parent_path = item.get("parentReference", {}).get("path", "")
        relative_parent = parent_path.split("root:", 1)[-1].strip("/")
        relative_path = "/".join(
            part for part in (relative_parent, item["name"]) if part
        )
        return {
            "id": item["id"],
            "name": item["name"],
            "type": "folder" if "folder" in item else "file",
            "size": item.get("size", 0),
            "modified_at": item.get("lastModifiedDateTime"),
            "web_url": item.get("webUrl"),
            "path": relative_path,
        }

    def list_files(self, folder_path: str = "") -> list[dict]:
        _, drive_id = self._site_and_drive_ids()
        endpoint = self._item_endpoint(drive_id, folder_path, "children")
        data = self._request("GET", endpoint).json()
        return [self._normalize_item(item) for item in data.get("value", [])]

    def search_files(self, query: str, limit: int = 20) -> list[dict]:
        query = query.strip()
        if not query:
            raise SharePointError("Search query is required")
        if not 1 <= limit <= MAX_SEARCH_RESULTS:
            raise SharePointError("Search limit must be between 1 and 200")
        _, drive_id = self._site_and_drive_ids()
        encoded_query = quote(query.replace("'", "''"), safe="")
        endpoint = f"/drives/{drive_id}/root/search(q='{encoded_query}')"
        data = self._request("GET", endpoint, params={"$top": limit}).json()
        return [
            self._normalize_item(item)
            for item in data.get("value", [])[:limit]
        ]

    def get_file_bytes(self, file_path: str, max_bytes: int) -> bytes:
        cleaned = self._clean_path(file_path)
        if not cleaned:
            raise SharePointError("File path is required")
        _, drive_id = self._site_and_drive_ids()
        endpoint = self._item_endpoint(drive_id, cleaned, "content")
        response = self._request("GET", endpoint)
        try:
            declared_size = int(response.headers.get("Content-Length", 0))
        except (TypeError, ValueError):
            declared_size = 0
        if declared_size > max_bytes or len(response.content) > max_bytes:
            raise SharePointError(f"File exceeds the {max_bytes}-byte limit")
        return response.content

    def read_text(self, file_path: str) -> dict:
        cleaned = self._clean_path(file_path)
        if Path(cleaned).suffix.lower() not in TEXT_EXTENSIONS:
            raise SharePointError("Unsupported text file type")
        content = self.get_file_bytes(cleaned, MAX_TEXT_BYTES)
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise SharePointError("Text file is not UTF-8 encoded") from exc
        return {"file": cleaned, "content": text, "bytes": len(content)}

    def _download_target(self, file_path: str) -> Path:
        cleaned = self._clean_path(file_path)
        if not cleaned:
            raise SharePointError("File path is required")
        root = self.config.download_dir.resolve()
        target = (root / Path(*PurePosixPath(cleaned).parts)).resolve()
        if not target.is_relative_to(root):
            raise SharePointError("Download path escapes the configured directory")
        return target

    def download_file(self, file_path: str) -> dict:
        content = self.get_file_bytes(file_path, MAX_DOWNLOAD_BYTES)
        target = self._download_target(file_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(content)
        return {
            "file": self._clean_path(file_path),
            "saved_path": str(target),
            "bytes": len(content),
        }

    def read_excel_range(
        self,
        file_path: str,
        sheet_name: str = "",
        cell_range: str = "",
    ) -> dict:
        cleaned = self._clean_path(file_path)
        if Path(cleaned).suffix.lower() != ".xlsx":
            raise SharePointError("Only .xlsx workbooks are supported")
        content = self.get_file_bytes(cleaned, MAX_EXCEL_BYTES)
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=True
            )
        except Exception as exc:
            raise SharePointError(f"Could not open Excel workbook: {exc}") from exc
        try:
            if not sheet_name:
                return {
                    "file": cleaned,
                    "worksheets": [
                        {
                            "name": sheet.title,
                            "rows": sheet.max_row,
                            "columns": sheet.max_column,
                        }
                        for sheet in workbook.worksheets
                    ],
                }
            if sheet_name not in workbook.sheetnames:
                raise SharePointError(f"Worksheet not found: {sheet_name}")
            sheet = workbook[sheet_name]
            selected_range = cell_range or (
                f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
            )
            try:
                min_col, min_row, max_col, max_row = range_boundaries(
                    selected_range
                )
            except ValueError as exc:
                raise SharePointError(
                    f"Invalid Excel range: {selected_range}"
                ) from exc
            bounds = (min_col, min_row, max_col, max_row)
            if any(value is None for value in bounds) or (
                min_col > max_col or min_row > max_row
            ):
                raise SharePointError(f"Invalid Excel range: {selected_range}")
            row_count = max_row - min_row + 1
            column_count = max_col - min_col + 1
            if (
                row_count > MAX_EXCEL_ROWS
                or row_count * column_count > MAX_EXCEL_CELLS
            ):
                raise SharePointError(
                    "Excel range exceeds the 200-row or 5,000-cell limit"
                )
            values = []
            for row in sheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
            ):
                values.append(
                    [
                        cell.value.isoformat()
                        if isinstance(cell.value, (date, datetime))
                        else cell.value
                        for cell in row
                    ]
                )
            return {
                "file": cleaned,
                "worksheet": sheet_name,
                "range": selected_range,
                "rows": row_count,
                "columns": column_count,
                "values": values,
            }
        finally:
            workbook.close()
